#!/usr/bin/env python3
"""Deterministically generate the synthetic test corpus.

    python make_corpus.py --out DIR

Every planted case from the build plan is here. Expected outcomes live in
corpus_expected.yaml. Bytes are reproducible: zip timestamps, document
properties, PDF ids and file mtimes are all pinned.
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import shutil
import struct
import zipfile
import zlib
from pathlib import Path

BASE = dt.datetime(2024, 1, 1, 12, 0, 0)
FIXED = (2024, 1, 1, 0, 0, 0)


def ts(days):
    return (BASE + dt.timedelta(days=days)).timestamp()


# ------------------------------------------------------------------- writers


def put(root, rel, text, days=0):
    path = root / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    os.utime(path, (ts(days), ts(days)))
    return path


def pin(path, days=0):
    """Rewrite a zip container with fixed entry timestamps and pin the mtime.

    openpyxl stamps the current time into docProps/core.xml on save, so the
    modified date is rewritten here as well.
    """
    modified = (BASE + dt.timedelta(days=days)).strftime("%Y-%m-%dT%H:%M:%SZ")
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            info = zipfile.ZipInfo(item.filename, FIXED)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = item.external_attr
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + modified.encode(), data)
            zout.writestr(info, data)
    tmp.replace(path)
    os.utime(path, (ts(days), ts(days)))


def make_docx(root, rel, title, paragraphs, days=0):
    from docx import Document

    path = root / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.author = "corpus"
    doc.core_properties.created = BASE
    doc.core_properties.modified = BASE + dt.timedelta(days=days)
    doc.core_properties.last_modified_by = "corpus"
    doc.add_heading(title, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(path)
    pin(path, days)


def make_xlsx(root, rel, sheet, rows, days=0):
    from openpyxl import Workbook

    path = root / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    wb.properties.creator = "corpus"
    wb.properties.created = BASE
    wb.properties.modified = BASE + dt.timedelta(days=days)
    wb.save(path)
    pin(path, days)


def make_pdf(root, rel, lines, days=0):
    from reportlab.pdfgen import canvas

    path = root / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), invariant=1, pagesize=(612, 792))
    y = 740
    for line in lines:
        c.drawString(54, y, line)
        y -= 18
        if y < 60:
            c.showPage()
            y = 740
    if not lines:
        c.rect(100, 100, 300, 300)  # a picture-only page: no extractable text
    c.showPage()
    c.save()
    os.utime(path, (ts(days), ts(days)))


def make_pptx(root, rel, slides, notes, days=0):
    from pptx import Presentation

    path = root / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.core_properties.author = "corpus"
    prs.core_properties.created = BASE
    prs.core_properties.modified = BASE + dt.timedelta(days=days)
    prs.core_properties.last_modified_by = "corpus"
    for i, (title, bullets) in enumerate(slides):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        tf = s.placeholders[1].text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b
        if i in notes:
            s.notes_slide.notes_text_frame.text = notes[i]
    prs.save(path)
    pin(path, days)


def tiny_png():
    def chunk(tag, data):
        c = struct.pack(">I", len(data)) + tag + data
        return c + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    raw = b"\x00\xff\x00\x00"
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


# ------------------------------------------------------------------- contents

ONBOARDING = """# Employee onboarding guide

Welcome to the company. During your first week you will receive a laptop, a building badge and
access to the shared calendar. Your manager schedules a short introduction with each member of
the team, and the people operations group runs an orientation session every Monday morning.

## First month

Complete the mandatory safety and compliance courses, set up direct deposit with payroll and
choose your benefits package. Buddies are assigned to every new hire to answer day to day
questions about tools, meetings and office etiquette.
"""

BRAND_VOICE = """# Brand voice guidelines

Our brand speaks plainly, warmly and with confidence. Avoid jargon, avoid exclamation marks and
prefer short sentences. Every campaign headline should promise one clear benefit for the reader.

## Tone by channel

Social posts may be lighter and use humour sparingly. Press releases stay formal. Customer emails
sound like a helpful colleague rather than a salesperson, and always end with a single clear
call to action.
"""

CODING = """---
title: Coding standards
owner: platform-team
reviewed: 2024-02-14
---
# Coding standards

All services are written in typed languages with strict compiler settings. Functions stay small,
names describe intent, and every public interface carries a short docstring.

## Reviews

Every change needs one approving review before merge. Reviewers check correctness first, then
naming, then style. Continuous integration must pass, and flaky tests are fixed or removed within
a week rather than retried.
"""

BATTERY = """Battery chemistry notes

Lithium iron phosphate cells trade some energy density for longer cycle life and better thermal
stability than nickel manganese cobalt cells. In our bench tests the cathode material dominated
degradation, while electrolyte additives slowed the growth of the solid electrolyte interphase.
Fast charging above two C accelerated lithium plating at low temperature.
"""

PRIVACY = [
    "Privacy notice",
    "We collect only the personal data needed to provide the service: your name, email address",
    "and billing details. We never sell personal data to third parties. You may request a copy",
    "of your data or ask us to delete it at any time by writing to the data protection officer.",
    "Cookies are used for sign in and for anonymous usage statistics. You can disable analytics",
    "cookies in the account settings without losing access to any feature.",
]

RUNBOOK = [
    "When a production alert fires, the on call engineer acknowledges it within five minutes.",
    "First check the dashboard for recent deploys, then the error budget, then upstream providers.",
    "If a rollback is needed, use the deploy tool and record the incident channel and timeline.",
    "After resolution schedule a blameless review within three working days and file follow up tasks.",
]

VENDOR = """# Vendor onboarding policy

Every new supplier must complete the vendor questionnaire before any purchase order is raised.
Procurement verifies tax registration, insurance certificates and banking details, and security
reviews any vendor that will handle company data. Approved vendors are added to the supplier
register and re-checked every twelve months.
"""

EXPENSE = """Expense policy

Employees may claim reasonable business expenses with an itemised receipt. Meals are reimbursed up
to a daily limit, flights must be booked in economy class, and hotel stays require manager
approval in advance. Claims are submitted within thirty days and paid with the next payroll run.
"""

SLIDES = [
    ("Platform roadmap", ["Where the platform is going in the next year", "Three themes: speed, safety, scale"]),
    ("Speed", ["Cut build times in half", "Introduce incremental deploys"]),
    ("Safety", ["Automated rollbacks for every service", "Quarterly disaster recovery drills"]),
    ("Scale", ["Multi region failover", "Capacity planning every quarter"]),
]
SLIDE_NOTES = {
    0: "Open with the customer outage story before introducing the three themes.",
    2: "Mention the failed drill from last quarter and what we changed afterwards.",
}

RELEASE = """# Release notes 4.2

This release improves startup time, fixes a crash when opening large projects and adds a dark theme.
The importer now handles files with unusual encodings, and the command palette remembers your most
recent commands. We also updated the bundled documentation and refreshed the sample projects so that
new users can follow the tutorials without extra setup. Known issue: the dark theme does not yet
cover the settings dialog. Thanks to everyone who reported problems during the beta period and
helped us verify the fixes on older hardware.
"""
RELEASE_EDITED = """# Release notes 4.2

This release improves startup time, fixes a crash when opening large projects and adds a dark theme.
The importer now handles files with unusual encodings, and the command palette remembers your most
recent commands. We also updated the bundled documentation and refreshed the sample projects so that
new users can follow the tutorials without extra setup. Known issue: the dark theme does not yet
cover the settings dialog. Thanks to everyone who reported problems during the beta.
"""

SPEC_V1 = """# Atlas design specification

Version: 1.0
Last updated: 2024-01-10

Atlas is a service that stores map tiles and serves them to the mobile apps. The first design uses
a single relational database for tile metadata and an object store for the tile images. Requests are
authenticated with static API keys, and cache invalidation is manual.
"""
SPEC_V2 = """# Atlas design specification

Version: 2.0
Last updated: 2024-02-10

Atlas is a service that stores map tiles and serves them to the mobile apps. The second design moves
tile metadata to a sharded key value store while the object store keeps the tile images. Requests are
authenticated with short lived tokens issued by the identity service. Cache invalidation is triggered
by publish events from the content pipeline, which removes the manual step.
"""
SPEC_V3 = """# Atlas design specification

Version: 3.0
Last updated: 2024-03-10

Atlas is a service that stores map tiles and serves them to the mobile apps and the partner website.
The third design keeps tile metadata in the sharded key value store and adds a regional edge cache in
front of the object store. Requests use short lived tokens, with per partner rate limits enforced at
the edge. Cache invalidation is driven by publish events, and stale tiles are refreshed in the
background after a configurable grace period.
"""

DEPLOY_PARAS = [
    "Deployments run through the continuous delivery pipeline. A merged change is built once, the artifact is signed, and the same artifact is promoted through staging and production without being rebuilt.",
    "Each environment has its own configuration store. Secrets are injected at start up from the vault and are never written to disk or to build logs. Operators rotate them on a fixed quarterly schedule.",
    "Rollouts use a canary stage that serves five percent of traffic for ten minutes. The pipeline watches error rate and latency, and it rolls back automatically if either crosses its threshold.",
    "Database migrations are applied before the new version starts and must be backward compatible with the previous release, so that a rollback never needs a schema change.",
    "After every production deploy the on call engineer verifies the health dashboard, confirms that the canary metrics returned to baseline and posts a short note in the release channel.",
    "Emergency fixes follow the same pipeline with a shortened approval step. Skipping the pipeline is only permitted during a declared incident and must be reviewed afterwards.",
]

REMOTE = """# Remote work policy

Employees may work remotely up to three days per week, provided their manager agrees and the team
keeps core hours free for meetings. Remote workers use the company VPN and keep their equipment
secure. The policy is reviewed every year by people operations.
"""
HANDBOOK = """# Employee handbook

This handbook summarises working conditions. Remote work is limited to two days per week and
requires manager approval. Core hours are ten to four. Holidays follow the national calendar and
unused leave carries over for one quarter. Questions go to people operations.
"""
SLA = """# Standard service level agreement

The platform is guaranteed 99.9% monthly uptime. Incidents of the highest severity receive a first
response within four hours. Credits are issued automatically when uptime falls below the guarantee.
Planned maintenance windows are announced one week ahead and do not count as downtime.
"""
SUPPORT = """# Support terms

Support is available on business days. Monthly uptime is guaranteed at 99.5% for standard plans.
Severe incidents receive a first response within eight hours, and customers on premium plans may
purchase faster response times. Maintenance is announced by email in advance.
"""

TRAVEL_OLD = """# Travel policy

Last updated: 2023-02-01

Business travel must be approved by your manager before booking. Flights are booked in economy
class through the corporate travel tool, and hotel rates are capped by city tier. Employees keep
receipts and submit them within thirty days. Rental cars require a valid licence and the basic
insurance bundle.

## Field stations

Field station travel: staff travelling to Antarctic research stations receive a per diem of 85 USD
for each day spent at the station, in addition to the standard allowances.
"""
TRAVEL_NEW = """# Travel policy

Last updated: 2024-06-10

Business travel must be approved by your manager before booking. Flights are booked in economy
class through the corporate travel tool, and hotel rates are capped by city tier. Employees keep
receipts and submit them within fourteen days. Rental cars require a valid licence and the basic
insurance bundle. Rail is preferred over flights for journeys under four hours.
"""

STATUS_OLD = """# Project status report

The migration is progressing. The database cutover is planned for next month and the data export
scripts are written. Testing of the import tooling has started. The main risk is the shortage of
test environments, which the infrastructure team is addressing.
"""
STATUS_FINAL = """# Project status report

The migration has completed its first phase. The database cutover is done and the old cluster is
read only. Import tooling passed its tests. Remaining work is decommissioning the old hardware and
updating the runbooks. Risks are closed and the infrastructure team has handed over.
"""

TICKET = """# PRJ-204 Homepage banner redesign

Design a new seasonal banner for the marketing homepage. Deliver three variants in light and dark
themes, with copy approved by the brand team, ready for the autumn campaign.
"""
BOARD = """# Sprint board

| ID | Title | Owner |
|----|-------|-------|
| PRJ-204 | Rotate database credentials | dana |
| PRJ-205 | Upgrade the build agents | lee |
"""


def build(root):
    root.mkdir(parents=True, exist_ok=True)
    # unique, unrelated documents
    put(root, "hr/onboarding_guide.md", ONBOARDING, 1)
    put(root, "marketing/brand_voice.md", BRAND_VOICE, 2)
    put(root, "engineering/coding_standards.md", CODING, 3)
    put(root, "research/battery_chemistry.txt", BATTERY, 4)
    make_pdf(root, "legal/privacy_notice.pdf", PRIVACY, 5)
    make_docx(root, "ops/incident_runbook.docx", "Incident runbook", RUNBOOK, 6)
    make_xlsx(root, "facilities/catering_schedule.xlsx", "Catering", [
        ["Day", "Caterer", "Meals"], ["Monday", "Green Fork", 120], ["Tuesday", "Soup Co", 95], ["Wednesday", "Green Fork", 130]], 7)
    # byte-identical pairs, one with conflicting names
    put(root, "forms/FRM-1042.md", VENDOR, 8)
    put(root, "policies/vendor_onboarding_policy.md", VENDOR, 9)
    put(root, "finance/expense_policy.txt", EXPENSE, 10)
    put(root, "finance/copy_of_expense_policy.txt", EXPENSE, 11)
    # deck with speaker notes and its PDF export
    make_pptx(root, "decks/platform_roadmap.pptx", SLIDES, SLIDE_NOTES, 12)
    lines = []
    for title, bullets in SLIDES:
        lines += [title] + bullets
    make_pdf(root, "decks/platform_roadmap.pdf", lines, 13)
    # near duplicate (not a revision chain)
    put(root, "notes/release_notes.md", RELEASE, 14)
    put(root, "notes/release_notes_edited.md", RELEASE_EDITED, 15)
    # revision chain with markers, text versions and mtimes in agreement
    put(root, "projects/atlas/design_spec.md", SPEC_V1, 10)
    put(root, "projects/atlas/design_spec_v2.md", SPEC_V2, 40)
    put(root, "projects/atlas/design_spec_v3.md", SPEC_V3, 70)
    # spreadsheet pair where only the text says one supersedes the other
    make_xlsx(root, "finance/headcount_plan_2024.xlsx", "Headcount", [
        ["Team", "Planned hires", "Budget"], ["Platform", 4, 520000], ["Support", 6, 410000], ["Design", 2, 230000],
        ["Notes", "Plan for fiscal year 2024"]], 20)
    make_xlsx(root, "finance/headcount_plan_2025.xlsx", "Headcount", [
        ["Team", "Planned hires", "Budget"], ["Platform", 5, 610000], ["Support", 3, 280000], ["Data", 3, 360000],
        ["Notes", "This plan supersedes the 2024 headcount plan"]], 30)
    # subsumed excerpt
    put(root, "wiki/deployment_overview.md", "# Deployment overview\n\n" + "\n\n".join(DEPLOY_PARAS) + "\n", 16)
    put(root, "wiki/deployment_overview_summary.md", "# Deployment overview summary\n\n" + "\n\n".join(DEPLOY_PARAS[1:4]) + "\n", 17)
    # planted contradictions
    put(root, "policy/remote_work_policy.md", REMOTE, 18)
    put(root, "hr/employee_handbook.md", HANDBOOK, 19)
    put(root, "ops/sla_standard.md", SLA, 20)
    put(root, "support/support_terms.md", SUPPORT, 21)
    # text-free pdf
    make_pdf(root, "scans/site_photo.pdf", [], 22)
    # superseded document holding unique content
    put(root, "policy/travel_policy.md", TRAVEL_OLD, 23)
    put(root, "policy/travel_policy_updated.md", TRAVEL_NEW, 90)
    # ambiguous marker, only mtime to go on
    put(root, "reports/status_report.md", STATUS_OLD, 100)
    put(root, "reports/status_report_final.md", STATUS_FINAL, 130)
    # identifier collision
    put(root, "engineering/PRJ-204_banner_redesign.md", TICKET, 24)
    put(root, "engineering/sprint_board.md", BOARD, 25)
    # not ingestible and unreadable
    (root / "misc").mkdir(parents=True, exist_ok=True)
    (root / "misc/diagram.png").write_bytes(tiny_png())
    (root / "misc/corrupt.docx").write_bytes(b"this is not a zip container " * 8)
    for rel in ("misc/diagram.png", "misc/corrupt.docx"):
        os.utime(root / rel, (ts(26), ts(26)))


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    out = Path(args.out)
    if out.exists() and any(out.iterdir()):
        shutil.rmtree(out)
    build(out)
    print(f"wrote {sum(1 for p in out.rglob('*') if p.is_file())} files to {out}")


if __name__ == "__main__":
    main()
